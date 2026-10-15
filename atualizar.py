#!/usr/bin/env python3
"""
Rotina diária do painel FMA dos 15 municípios do PMIF.

Para cada município (municipios.json): baixa umidade e chuva horárias do modelo ECMWF (via Open-Meteo) no ponto de
cálculo do município, atualiza o histórico (dados/<codigo_ibge>_diario.csv), recalcula a FMA sobre todo o histórico e
gera o painel (docs/index.html e docs/dados.json).

Rodar:  python atualizar.py
"""
import json
import math
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pandas as pd

from fma import MODELO_PADRAO, baixar_openmeteo_horario, calcular_fma, diario_de_openmeteo

RAIZ = Path(__file__).parent
DADOS = RAIZ / "dados"
DOCS = RAIZ / "docs"
DIAS_INICIAL = 90        # primeira execução: baixa 90 dias para aquecer o acumulado
DIAS_ATUALIZACAO = 10    # execuções seguintes: reprocessa os últimos 10 dias (corrige dados atrasados)
DIAS_NO_PAINEL = 60
DIAS_PREVISAO = 3        # a previsão só aparece se validacao_previsao.json existir (ver validar_previsao.py)
FUSO = timezone(timedelta(hours=-3))  # Brasília


def dist_km(lat1, lon1, lat2, lon2):
    p1, p2 = math.radians(lat1), math.radians(lat2)
    h = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(lon2 - lon1) / 2) ** 2
    return 2 * 6371.0 * math.asin(math.sqrt(h))


def estacao_mais_proxima(lat, lon):
    """Estação INMET (das usadas na validação) mais próxima do município; só para contexto no painel."""
    arq = RAIZ / "dados_estacoes" / "estacoes_meta.json"
    if not arq.exists():
        return None
    meta = json.loads(arq.read_text(encoding="utf-8"))
    cod = min(meta, key=lambda c: dist_km(lat, lon, meta[c]["latitude"], meta[c]["longitude"]))
    return {"codigo": cod, "nome": meta[cod]["nome"], "km": round(dist_km(lat, lon, meta[cod]["latitude"], meta[cod]["longitude"]))}


def atualizar_municipio(m, hoje, com_previsao=False):
    cod = m["codigo_ibge"]
    arq = DADOS / f"{cod}_diario.csv"

    if arq.exists():
        hist = pd.read_csv(arq, parse_dates=["data"]).set_index("data")
        dias = max(DIAS_ATUALIZACAO, (hoje - hist.index.max().date()).days + 2)  # cobre falhas de vários dias
    else:
        hist = pd.DataFrame(columns=["umidade", "chuva"])
        hist.index.name = "data"
        dias = DIAS_INICIAL

    h = baixar_openmeteo_horario(m["latitude"], m["longitude"], dias + 1,
                                 dias_previsao=DIAS_PREVISAO if com_previsao else 0)
    novo = diario_de_openmeteo(h, datetime.now(FUSO))
    if not novo.empty:
        novo["data"] = pd.to_datetime(novo["data"])
        novo = novo[novo["data"].dt.date <= hoje].set_index("data")
        # hoje, antes da leitura das 13h, ainda não há umidade: não mostra o dia como "Sem dado"
        novo = novo[~((novo.index.date == hoje) & novo["umidade"].isna())]
        # valores novos têm prioridade; lacunas são preenchidas com o que já estava guardado
        hist = novo.combine_first(hist) if len(hist) else novo

    if hist.empty:
        raise RuntimeError(f"{m['nome']}: nenhum dado disponível")

    hist = hist.sort_index()
    hist = hist.reindex(pd.date_range(hist.index.min(), hist.index.max(), name="data"))  # dias sem dado aparecem
    hist[["umidade", "chuva"]].to_csv(arq, date_format="%Y-%m-%d")

    res = calcular_fma(hist.reset_index())
    res.to_csv(DADOS / f"{cod}_fma.csv", index=False, date_format="%Y-%m-%d")   # histórico: só o que já aconteceu
    res["previsao"] = False

    if com_previsao:
        # dias à frente do último dia observado (até DIAS_PREVISAO); a FMA continua do acumulado desse último dia
        fut = diario_de_openmeteo(h)
        fut["data"] = pd.to_datetime(fut["data"])
        fut = fut.set_index("data")
        ultimo = hist.index.max()
        fut = fut[(fut.index > ultimo) & (fut.index <= ultimo + timedelta(days=DIAS_PREVISAO))]   # no máximo 3 dias além do último dia observado
        if not fut.empty:
            tudo = pd.concat([hist[["umidade", "chuva"]], fut[["umidade", "chuva"]]])
            tudo = tudo.reindex(pd.date_range(tudo.index.min(), tudo.index.max(), name="data"))
            r2 = calcular_fma(tudo.reset_index())
            r2["previsao"] = r2["data"] > ultimo
            res = r2
    return res


def serie_para_painel(res):
    linhas = []
    n_prev = int(res["previsao"].sum()) if "previsao" in res else 0
    for _, r in res.tail(DIAS_NO_PAINEL + n_prev).iterrows():
        linhas.append({
            "data": r["data"].strftime("%Y-%m-%d"),
            "umidade": None if pd.isna(r["umidade"]) else round(float(r["umidade"]), 1),
            "chuva": None if pd.isna(r["chuva"]) else round(float(r["chuva"]), 1),
            "fma": None if pd.isna(r["fma"]) else round(float(r["fma"]), 2),
            "classe": r["classe"],
            "aviso": r["aviso"],
            "previsao": bool(r.get("previsao", False)),
        })
    return linhas


def gerar_historico_csv(municipios):
    """Um único CSV com o histórico completo de todos os municípios (docs/historico_fma.csv), para baixar pelo painel.
    Formato para Excel em português: separador ';', decimal ',' e acentos corretos."""
    partes = []
    for m in municipios:
        arq = DADOS / f"{m['codigo_ibge']}_fma.csv"
        if not arq.exists():
            continue
        d = pd.read_csv(arq)
        d.insert(0, "municipio", m["nome"])
        d.insert(0, "codigo_ibge", m["codigo_ibge"])
        partes.append(d)
    if not partes:
        return
    tudo = pd.concat(partes, ignore_index=True)
    tudo["fma"] = tudo["fma"].round(2)
    tudo["chuva"] = tudo["chuva"].round(1)
    tudo["umidade"] = tudo["umidade"].round(1)
    tudo = tudo.rename(columns={"umidade": "umidade_13h_pct", "chuva": "chuva_24h_mm", "fma": "fma"})
    tudo.to_csv(DOCS / "historico_fma.csv", index=False, sep=";", decimal=",", encoding="utf-8-sig")


def gerar_historico_xlsx(municipios):
    """Excel com uma aba por município (docs/historico_fma.xlsx) e uma aba "Leia-me" com as regras e as fontes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cores = {"Nulo": "DCE9F5", "Pequeno": "D5EBDC", "Médio": "FBF0B3", "Alto": "FBD9B8", "Muito alto": "F3C0C0", "Sem dado": "E5E7E9"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Leia-me"
    linhas = [
        ("Histórico da FMA (Fórmula de Monte Alegre) - municípios do PMIF", True),
        (f"Gerado em {datetime.now(FUSO).strftime('%d/%m/%Y às %H:%M')} (Brasília). Uma aba por município.", False),
        ("", False),
        ("Dados: umidade relativa às 13h e chuva (das 14h do dia anterior às 13h do dia) do modelo ECMWF IFS 0,25° via Open-Meteo, no ponto de cálculo de cada município.", False),
        ("São estimativas de modelo, não medições de estação.", False),
        ("Chuva até 2,4 mm: sem efeito; 2,5 a 4,9: acumulado x 0,7; 5,0 a 9,9: x 0,4; 10,0 a 12,9: x 0,2; acima de 12,9: zera.", False),
        ("Classes: Nulo (até 1,0), Pequeno (1,1 a 3,0), Médio (3,1 a 8,0), Alto (8,1 a 20,0), Muito alto (acima de 20,0).", False),
        ("Referência dos limites das classes: SOARES, R. V. Determinação de um índice de perigo de incêndio para a região centro paranaense, Brasil. Turrialba, Costa Rica, CATIE/IICA. 72 p. (Tese de Mestrado). 1972.", False),
        ("Dias sem umidade às 13h aparecem como Sem dado. A coluna aviso indica chuva ausente (assumida 0 mm).", False),
        ("A FMA é acumulativa: no início da série o valor pode estar subestimado.", False),
    ]
    for i, (t, neg) in enumerate(linhas, 1):
        ws.cell(row=i, column=1, value=t).font = Font(bold=neg, size=13 if neg else 11)
    ws.column_dimensions["A"].width = 120

    cab = ["Data", "Umidade 13h (%)", "Chuva 24h (mm)", "FMA", "Classe", "Aviso"]
    larg = [12, 17, 16, 10, 13, 34]
    for m in municipios:
        arq = DADOS / f"{m['codigo_ibge']}_fma.csv"
        if not arq.exists():
            continue
        d = pd.read_csv(arq, parse_dates=["data"]).fillna({"aviso": ""})
        aba = wb.create_sheet(m["nome"][:31])
        aba.append(cab)
        for c in range(1, len(cab) + 1):
            cel = aba.cell(row=1, column=c)
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="2F5D50")
            cel.alignment = Alignment(horizontal="center")
            aba.column_dimensions[get_column_letter(c)].width = larg[c - 1]
        for _, r in d.iterrows():
            aba.append([r["data"].to_pydatetime(),
                        None if pd.isna(r["umidade"]) else round(float(r["umidade"]), 1),
                        None if pd.isna(r["chuva"]) else round(float(r["chuva"]), 1),
                        None if pd.isna(r["fma"]) else round(float(r["fma"]), 2),
                        r["classe"], r["aviso"]])
            n = aba.max_row
            aba.cell(row=n, column=1).number_format = "dd/mm/yyyy"
            aba.cell(row=n, column=2).number_format = "0.0"
            aba.cell(row=n, column=3).number_format = "0.0"
            aba.cell(row=n, column=4).number_format = "0.00"
            cor = cores.get(r["classe"])
            if cor:
                aba.cell(row=n, column=5).fill = PatternFill("solid", fgColor=cor)
        aba.freeze_panes = "A2"
    wb.save(DOCS / "historico_fma.xlsx")


def gerar_painel(municipios, resultados, erros, validacao=None):
    gerar_historico_csv(municipios)
    try:
        gerar_historico_xlsx(municipios)
    except Exception as e:   # o Excel é um extra: se falhar, o painel segue
        print(f"Aviso: não foi possível gerar o Excel: {e}", file=sys.stderr)
    agora = datetime.now(FUSO)
    saida = {
        "atualizado": agora.strftime("%d/%m/%Y às %H:%M (Brasília)"),
        "modelo": "ECMWF IFS 0,25° (via Open-Meteo)",
        "modelo_id": MODELO_PADRAO,
        "erros": erros,
        "previsao": {"dias": DIAS_PREVISAO, "validacao": validacao} if validacao else None,
        "municipios": [
            {"codigo_ibge": m["codigo_ibge"], "nome": m["nome"], "latitude": m["latitude"], "longitude": m["longitude"],
             "estacao_proxima": estacao_mais_proxima(m["latitude"], m["longitude"]),
             "serie": serie_para_painel(resultados[m["codigo_ibge"]])}
            for m in municipios if m["codigo_ibge"] in resultados
        ],
    }
    (DOCS / "dados.json").write_text(json.dumps(saida, ensure_ascii=False, indent=1), encoding="utf-8")
    # o painel leva também os contornos (só no HTML, para o dados.json continuar leve)
    com_mapa = dict(saida)
    com_mapa["contornos"] = {m["codigo_ibge"]: m["poligonos"] for m in municipios}
    modelo = (RAIZ / "painel_template.html").read_text(encoding="utf-8")
    html = modelo.replace("__DADOS__", json.dumps(com_mapa, ensure_ascii=False, separators=(",", ":")))
    (DOCS / "index.html").write_text(html, encoding="utf-8")


def main():
    DADOS.mkdir(exist_ok=True)
    DOCS.mkdir(exist_ok=True)
    municipios = json.loads((RAIZ / "municipios.json").read_text(encoding="utf-8"))
    hoje = datetime.now(FUSO).date()
    arq_val = RAIZ / "validacao_previsao.json"
    validacao = json.loads(arq_val.read_text(encoding="utf-8")) if arq_val.exists() else None
    if validacao is None:
        print("Sem validacao_previsao.json: a previsão fica desligada (rode a validação para ligar).")

    resultados, erros = {}, {}
    pendentes = list(municipios)
    for rodada in range(1, 4):   # a API gratuita às vezes demora: quem falhar é tentado de novo, até 3 rodadas
        falharam = []
        for m in pendentes:
            cod = m["codigo_ibge"]
            try:
                resultados[cod] = atualizar_municipio(m, hoje, com_previsao=validacao is not None)
                erros.pop(cod, None)
                ult = resultados[cod][~resultados[cod]["previsao"]].dropna(subset=["fma"]).tail(1)
                if not ult.empty:
                    r = ult.iloc[0]
                    print(f"{m['nome']}: FMA {r['fma']:.2f} ({r['classe']}) em {r['data']:%d/%m/%Y}")
            except Exception as e:
                print(f"ERRO em {m['nome']} (rodada {rodada}): {e}", file=sys.stderr)
                erros[cod] = {"codigo": cod, "nome": m["nome"], "erro": str(e)}
                falharam.append(m)
            time.sleep(1)   # gentileza com a API gratuita
        pendentes = falharam
        if not pendentes or rodada == 3:
            break
        print(f"{len(pendentes)} município(s) falharam; nova tentativa em 90 s.", file=sys.stderr)
        time.sleep(90)

    for m in pendentes:   # falhou mesmo assim: mostra o último resultado guardado, se existir
        arq = DADOS / f"{m['codigo_ibge']}_fma.csv"
        if arq.exists():
            resultados[m["codigo_ibge"]] = pd.read_csv(arq, parse_dates=["data"]).fillna({"aviso": ""})
            resultados[m["codigo_ibge"]]["previsao"] = False
    erros = list(erros.values())

    gerar_painel(municipios, resultados, erros, validacao)
    if erros and not resultados:
        sys.exit(1)


if __name__ == "__main__":
    main()
